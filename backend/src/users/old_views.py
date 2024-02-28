from typing import Any
import jwt, datetime
from mealmanager.settings._base import JWT_SALT
from django.core.exceptions import ObjectDoesNotExist
from openpyxl import Workbook, load_workbook
from rest_framework.decorators import api_view, renderer_classes
from rest_framework.views import APIView
from rest_framework.exceptions import AuthenticationFailed
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, FileUploadParser
from rest_framework.response import Response
from rest_framework import generics, permissions, status
from .serializers import (
    BatchUploadUserProfileSerializer,
    UserSerializer,
    UserProfileSerializer,
    AvatarSerializer,
)
from .models import User, UserProfile, EmployeeBatchUpload
from .auth_service import user_auth


# Create your views here.
class RegisterView(APIView):
    def post(self, request):
        serializer = UserSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)


class LoginView(APIView):
    def post(self, request):
        email = request.data["email"]
        password = request.data["password"]

        user = User.objects.filter(email=email).first()
        if user is None:
            # raise AuthenticationFailed("User not found!")
            return Response(
                {"invalid_user_error": "User not found!"},
                status=status.HTTP_403_FORBIDDEN,
            )

        if not user.check_password(password):
            # raise AuthenticationFailed("Incorrect password!")
            return Response(
                {"password_error": "Incorrect password!"},
                status=status.HTTP_403_FORBIDDEN,
            )

        payload = {
            "id": user.id,
            "exp": datetime.datetime.utcnow() + datetime.timedelta(days=365 * 10),
            "iat": datetime.datetime.utcnow(),
        }
        token = jwt.encode(payload, JWT_SALT, algorithm="HS256")

        response = Response()

        response.set_cookie(key="jwt", value=token, httponly=True)
        response.data = {"jwt": token}

        return response


class AuthUserView(APIView):
    def get(self, request, pk=None):
        payload = user_auth(request)
        if payload.get("auth_error", None):
            return Response(payload, status=status.HTTP_403_FORBIDDEN)
        user = User.objects.filter(id=payload["id"]).first()
        try:
            user_serializer = UserSerializer(user)
            profile = UserProfile.objects.select_related("user").get(user=user.pk)
            profile_serialiser = UserProfileSerializer(profile)
            return Response(data=profile_serialiser.data)
        except UserProfile.DoesNotExist:
            return Response(data=user_serializer.data, status=status.HTTP_200_OK)


class LogoutView(APIView):
    def post(self, request):
        response = Response()
        response.delete_cookie("jwt")
        response.data = {"message": "success"}
        return response


class UserProfileView(APIView):
    def get(self, request, pk=None):
        payload = user_auth(request)
        if payload.get("auth_error", None):
            return Response(payload, status=status.HTTP_403_FORBIDDEN)
        try:
            if pk is not None:
                profile = UserProfile.objects.select_related("user").get(user=pk)
                profile_serialiser = UserProfileSerializer(profile)
                return Response(data=profile_serialiser.data)
            else:
                profiles = UserProfile.objects.select_related("user").all()
                profiles_serializer = UserProfileSerializer(profiles, many=True)
                return Response(data=profiles_serializer.data)
        except UserProfile.DoesNotExist as e:
            return Response(
                data={"no_profile_error": e.args[0]}, status=status.HTTP_400_BAD_REQUEST
            )

    def post(self, request):
        payload = user_auth(request)
        if payload.get("auth_error", None):
            return Response(payload, status=status.HTTP_403_FORBIDDEN)
        auth_user = User.objects.get(id=payload["id"])
        user_data = {
            "first_name": request.data.pop("first_name"),
            "last_name": request.data.pop("last_name"),
            "email": auth_user.email,
            "username": auth_user.username,
            "password": auth_user.password,
        }
        ###...Update related User before creating Profile
        user_serializer = UserSerializer(instance=auth_user, data=user_data)
        profile_data = {**request.data}
        print("#####Profile Data:", profile_data)
        profile_data["user"] = {
            "id": auth_user.id,
            "email": "...",
            "username": "...",
            "password": "...",
        }
        profile_data["user_id"] = auth_user.id
        profile_data["reader_uid"] = ""
        print("%%%%%%:", profile_data["user_id"])

        profile_serialiser = UserProfileSerializer(data=profile_data)
        if profile_serialiser.is_valid() and user_serializer.is_valid():
            # ....this suite is a candidate for DB Transaction
            user_serializer.save()
            profile_serialiser.save()
            print("****CreatedData****:", profile_serialiser.data)
            created_data = {
                **profile_serialiser.data,
            }
            return Response(data=created_data, status=status.HTTP_200_OK)
        else:
            profile_serialiser.is_valid()
            user_serializer.is_valid()
            errors = [profile_serialiser.errors, user_serializer.errors]
            return Response(data=errors, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    def patch(self, request, format=None):
        payload = user_auth(request)
        if payload.get("auth_error", None):
            return Response(payload, status=status.HTTP_403_FORBIDDEN)
        auth_user = User.objects.get(id=payload["id"])
        user_data = {
            "first_name": request.data.pop("first_name", auth_user.first_name),
            "last_name": request.data.pop("last_name", auth_user.last_name),
            "email": auth_user.email,
            "username": auth_user.username,
            "password": auth_user.password,
        }
        user_serialiser = UserSerializer(instance=auth_user, data=user_data)

        profile = UserProfile.objects.get(user=auth_user.pk)
        profile_data = {**request.data}
        profile_data["user"] = {
            "id": auth_user.id,
            "email": "...",
            "username": "...",
            "password": "...",
        }
        profile_data["user_id"] = auth_user.id
        profile_data["reader_uid"] = profile.reader_uid
        print("****MergeData****:", profile_data)
        profile_serialiser = UserProfileSerializer(instance=profile, data=profile_data)
        if profile_serialiser.is_valid() and user_serialiser.is_valid():
            user_serialiser.save()
            profile_serialiser.save()
            print("****UpdateData****:", profile_serialiser.data)
            updated_data = {
                **profile_serialiser.data,
            }
            return Response(data=updated_data, status=status.HTTP_200_OK)
        else:
            print("Profile PATCH ERROR: ", profile_serialiser.errors)
            return Response(
                data=profile_serialiser.errors,
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )


class UploadProfileImageView(APIView):
    parser_classes = [FormParser, MultiPartParser, FileUploadParser]

    def patch(self, request, format=None):
        payload = user_auth(request)
        if payload.get("auth_error", None):
            return Response(payload, status=status.HTTP_403_FORBIDDEN)
        user = User.objects.get(id=payload["id"])
        profile = UserProfile.objects.get(user=user.pk)
        serialiser = AvatarSerializer(instance=profile, data=request.data)
        if serialiser.is_valid():
            serialiser.save()
            return Response(data=serialiser.data, status=status.HTTP_200_OK)
        else:
            return Response(data=serialiser.errors, status=status.HTTP_400_BAD_REQUEST)


# @api_view(["POST"])
def bulk_create(request):
    # payload: dict[str, str] | Any = user_auth(request)
    # if payload.get("auth_error", None):
    # return Response(payload, status=status.HTTP_403_FORBIDDEN)
    # _user: Any = User.objects.filter(id=payload["id"]).first()
    # if not _user.is_superuser or not _user.is_staff:
    # return Response(
    # data={"error": "User is not an Admin!"}, status=status.HTTP_401_UNAUTHORIZED
    # )
    try:
        batch_file: Any = EmployeeBatchUpload.objects.get(
            batch_file__icontains="employee_batch"
        ).batch_file
    except EmployeeBatchUpload.DoesNotExist:
        return Response(
            data={"error": "No batch file found!"}, status=status.HTTP_400_BAD_REQUEST
        )
    wb: Workbook = load_workbook(filename=batch_file)
    ws_user = wb.worksheets[0]
    ws_profile = wb.worksheets[1]
    imported_user_counter = 0
    skipped_user_counter = 0
    imported_profile_counter = 0
    skipped_profile_counter = 0
    ws_user_columns = [
        "first_name",
        "last_name",
        "is_staff",
        "is_active",
        "username",
        "email",
        "password",
    ]
    ws_profile_columns = [
        "reader_uid",
        "meal_category",
        "department",
        "user",
        "user_id",
    ]
    user_schema_column = {
        "id": "",
        "email": "...",
        "username": "...",
        "password": "...",
    }

    user_rows = ws_user.iter_rows(min_row=2)
    profile_rows = ws_profile.iter_rows(min_row=2)
    for index, user_row in enumerate(user_rows, start=1):
        user_row_values = [cell.value for cell in user_row[1:]]
        user_row_dict = dict(zip(ws_user_columns, user_row_values))
        print("@@@@ Total User From excel: ", len(user_row_dict.keys()))
        user_serializer = UserSerializer(data=user_row_dict)
        if user_serializer.is_valid():
            print("@@@@ VALID USER: ")
            user_serializer.save()
            print("@@@@ VALID USER DATA: ", user_serializer.data)
            imported_user_counter += 1
        else:
            print("@@@@ NOT VALID USER:", user_serializer.errors)
            skipped_user_counter += 1

    for index, profile_row in enumerate(profile_rows, start=1):
        profile_row_values = [cell.value for cell in profile_row[1:]]
        print("#### PROFILE ROLE VALUES: ", profile_row_values)
        email = profile_row_values.pop()
        try:
            user = User.objects.get(email=email)

            u_serializer = UserSerializer(user)
            user_schema_column["id"] = u_serializer.data["id"]
            profile_row_values.append(user_schema_column)
            profile_row_values.append(u_serializer.data["id"])
            print("$$$$ ROW VALUES:", profile_row_values)

            profile_row_dict = dict(zip(ws_profile_columns, profile_row_values))
            print("%%%%%%%%- ROW DICT:", profile_row_dict)

            p_serializer = BatchUploadUserProfileSerializer(data=profile_row_dict)
            if p_serializer.is_valid():
                print("@@@@  VALID PROFILE:")
                p_serializer.save()
                imported_profile_counter += 1
                print("####  imported PROFILE:", imported_profile_counter)
            else:
                print("@@@@ NOT VALID PROFILE:", p_serializer.errors)
                skipped_profile_counter += 1
                print("#### skipped PROFILE:", skipped_profile_counter)
        except User.DoesNotExist:
            error = {"error": f"user with email: {email} was not found!"}
            return Response(
                data=error,
                status=status.HTTP_400_BAD_REQUEST,
            )
        except Exception as e:
            print("PROFILE EXCEPTION @310:", e)
            return Response(
                data={"error": e.args[0]},
                status=status.HTTP_400_BAD_REQUEST,
            )
    if skipped_user_counter > 1 or skipped_profile_counter > 0:
        data = {
            "total_user_failed": skipped_user_counter,
            "total_profile_failed": skipped_profile_counter,
        }
        print("Skiped User Count: ", skipped_user_counter)
        return Response(
            data=data,
            status=status.HTTP_400_BAD_REQUEST,
        )
    if imported_user_counter == imported_profile_counter and (
        imported_user_counter > 0 and imported_profile_counter > 0
    ):
        data = (
            {
                "total_uploaded_users": imported_user_counter,
                "total_uploaded_profiles": imported_profile_counter,
            },
        )
        return Response(
            data=data,
            status=status.HTTP_200_OK,
        )
    return Response(
        data={
            "total_uploaded_users": imported_user_counter,
            "total_uploaded_profiles": imported_profile_counter,
        },
    )
